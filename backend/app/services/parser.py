"""Typed doc parser — PDF / Word / Markdown / TXT -> ParsedDocument.

Each parse_* function returns the same shape so callers don't care about file type.
All heavy I/O happens *synchronously* (FastAPI endpoint wraps in threadpool).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from app.core.exceptions import UnsupportedFileTypeError

ChunkType = Literal["heading", "text", "code", "list", "table"]

SUPPORTED_TYPES: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
    ".text": "txt",
}


@dataclass
class Block:
    """A typed block inside a document."""
    text: str
    type: ChunkType = "text"
    level: int | None = None          # For headings: 1..6


@dataclass
class ParsedDocument:
    text: str
    blocks: list[Block] = field(default_factory=list)
    word_count: int = 0
    metadata: dict = field(default_factory=dict)


# -------- public entrypoints --------

def parse_document(file_path: str | Path) -> ParsedDocument:
    path = Path(file_path)
    ext = path.suffix.lower()
    file_type = SUPPORTED_TYPES.get(ext)
    if file_type is None:
        raise UnsupportedFileTypeError(
            message=f"Unsupported file type: {ext}",
            detail=f"Supported: {', '.join(SUPPORTED_TYPES.keys())}",
        )
    if file_type == "pdf":
        return _parse_pdf(path)
    if file_type == "docx":
        return _parse_docx(path)
    if file_type == "md":
        return _parse_markdown(path)
    return _parse_txt(path)


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    ft = SUPPORTED_TYPES.get(ext)
    if ft is None:
        raise UnsupportedFileTypeError(
            message=f"Unsupported file type: {ext}",
            detail=f"Supported: {', '.join(SUPPORTED_TYPES.keys())}",
        )
    return ft


# -------- PDF --------

def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    parts: list[str] = []
    metadata: dict = {}
    try:
        info = reader.metadata or {}
        metadata = {
            "title": getattr(info, "title", None),
            "author": getattr(info, "author", None),
            "page_count": len(reader.pages),
        }
    except Exception:
        pass

    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if not txt.strip():
            continue
        parts.append(txt)
        for para in re.split(r"\n\s*\n+", txt.strip()):
            para = para.strip()
            if not para:
                continue
            t: ChunkType = "heading" if _looks_like_heading(para) else "text"
            blocks.append(Block(text=para, type=t))

    text = "\n\n".join(parts)
    return ParsedDocument(text=text, blocks=blocks, word_count=len(text.split()), metadata=metadata)


# -------- DOCX --------

def _parse_docx(path: Path) -> ParsedDocument:
    import docx

    d = docx.Document(str(path))
    blocks: list[Block] = []
    parts: list[str] = []

    for para in d.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            m = re.search(r"\d+", style)
            lvl = int(m.group()) if m else 1
            blocks.append(Block(text=txt, type="heading", level=lvl))
        elif style.startswith("code") or _looks_like_code(txt):
            blocks.append(Block(text=txt, type="code"))
        else:
            blocks.append(Block(text=txt, type="text"))
        parts.append(txt)

    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            t = "\n".join(rows)
            blocks.append(Block(text=t, type="table"))
            parts.append(t)

    text = "\n\n".join(parts)
    return ParsedDocument(text=text, blocks=blocks, word_count=len(text.split()))


# -------- MARKDOWN --------

def _parse_markdown(path: Path) -> ParsedDocument:
    md_text = path.read_text(encoding="utf-8")

    # Use markdown-it for accuracy — especially heading level.
    try:
        from markdown_it import MarkdownIt
        md = MarkdownIt("commonmark", {"html": False})
        tokens = md.parse(md_text)

        blocks: list[Block] = []
        parts: list[str] = []
        buf: list[str] = []
        current_heading_level: int | None = None
        current_fence_lang: str | None = None

        def _flush_para() -> None:
            if buf:
                t = "\n".join(buf).strip()
                if t:
                    lvl = current_heading_level
                    btype: ChunkType = "heading" if lvl is not None else "text"
                    blocks.append(Block(text=t, type=btype, level=lvl))
                    parts.append(t)
                buf.clear()

        for tok in tokens:
            if tok.type == "heading_open":
                _flush_para()
                # markdown-it-py: heading level is on tok.tag ("h1".."h6"), not hLevel attr.
                tag = getattr(tok, "tag", "") or ""
                m = re.match(r"h([1-6])", tag)
                current_heading_level = int(m.group(1)) if m else 1
                buf.clear()
            elif tok.type == "heading_close":
                _flush_para()
                current_heading_level = None
            elif tok.type == "fence":
                _flush_para()
                if tok.content.strip():
                    blocks.append(Block(text=tok.content.strip(), type="code"))
                    parts.append(tok.content.strip())
                current_fence_lang = None
            elif tok.type == "paragraph_open":
                _flush_para()
            elif tok.type == "inline":
                buf.append(tok.content)
            elif tok.type == "paragraph_close":
                _flush_para()

        _flush_para()

        # Fallback: no heading captured? Use simple ATX split.
        if not any(b.type == "heading" for b in blocks):
            blocks = _split_atx(md_text)
            parts = [b.text for b in blocks]

        text = "\n\n".join(parts)
        return ParsedDocument(text=text, blocks=blocks, word_count=len(text.split()))

    except Exception:
        # markdown-it missing or failed — simple regex fallback.
        blocks = _split_atx(md_text)
        text = "\n\n".join(b.text for b in blocks)
        return ParsedDocument(text=text, blocks=blocks, word_count=len(text.split()))


def _split_atx(md_text: str) -> list[Block]:
    """ATX-heading-aware fallback splitter (used when markdown-it unavailable)."""
    blocks: list[Block] = []
    buf: list[str] = []
    in_fence = False

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if buf:
                blocks.append(Block(text="\n".join(buf).strip(), type="text"))
                buf = []
            in_fence = not in_fence
            continue

        if not in_fence:
            m = re.match(r"^(#{1,6})\s+(.*)", line)
            if m:
                if buf:
                    blocks.append(Block(text="\n".join(buf).strip(), type="text"))
                    buf = []
                blocks.append(Block(text=m.group(2).strip(), type="heading", level=len(m.group(1))))
                continue

        buf.append(line)

    if buf:
        text = "\n".join(buf).strip()
        if text:
            blocks.append(Block(text=text, type="text"))
    return blocks


# -------- TXT --------

def _parse_txt(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8")
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n+", text.strip()):
        para = para.strip()
        if not para:
            continue
        t: ChunkType = "heading" if _looks_like_heading(para) else "text"
        blocks.append(Block(text=para, type=t))
    return ParsedDocument(text=text, blocks=blocks, word_count=len(text.split()))


# -------- helpers --------

def _looks_like_heading(text: str) -> bool:
    if len(text) > 80 or len(text) < 3:
        return False
    if text[-1] in ".!?。！？,，;；":
        return False
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return False
    upper_count = sum(1 for c in letters if c.isupper())
    ascii_count = sum(1 for c in text if c.isascii() and c.isalpha())
    if ascii_count == 0:
        return True
    return upper_count / max(ascii_count, 1) > 0.5


def _looks_like_code(text: str) -> bool:
    if any(kw in text for kw in ("def ", "class ", "import ", "function ", "const ", "public ")):
        return True
    return text.count("{") + text.count("}") + text.count("=") > text.count(" ") * 0.3


def compute_file_hash(file_path: Path) -> str:
    import hashlib
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()
